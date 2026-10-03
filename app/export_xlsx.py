"""Ekspor Excel: 3 sheet mengikuti FORMAT_SPJ_KALKULASI_PERJALANAN_DINAS_DPRD.xlsx
(DP RIIL, SP TANGGUNG JAWAB, KALKULASI) dengan rumus hidup."""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.worksheet.properties import PageSetupProperties

from .calc import (KATEGORI_BIAYA, PESAWAT, Result, Settings, Trip, hitung,
                   nama_provinsi, rentang_tgl, tgl_id)

THIN = Side(style="thin")
NUM_ACC = '_(* #,##0_);_(* \\(#,##0\\);_(* "-"??_);_(@_)'
NUM_RP = '_([$Rp-421]* #,##0.00_);_([$Rp-421]* \\(#,##0.00\\);_([$Rp-421]* "-"??_);_(@_)'


def F(sz=11, b=False, u=False):
    return Font(name="Arial", size=sz, bold=b, underline="single" if u else None)


def A(h=None, v="center", wrap=False):
    return Alignment(horizontal=h, vertical=v, wrap_text=wrap)


def setc(ws, ref, value=None, font=None, align=None, fmt=None):
    c = ws[ref]
    if value is not None:
        c.value = value
    if font:
        c.font = font
    if align:
        c.alignment = align
    if fmt:
        c.number_format = fmt
    return c


def box(ws, first_row, last_row, c1, c2, inner_top=None):
    """Bingkai luar area c1..c2 (indeks kolom)."""
    for r in range(first_row, last_row + 1):
        for c in range(c1, c2 + 1):
            cell = ws.cell(r, c)
            b = cell.border
            cell.border = Border(
                left=THIN if c == c1 else b.left, right=THIN if c == c2 else b.right,
                top=THIN if r == first_row else b.top, bottom=THIN if r == last_row else b.bottom)


def a4(ws, fit_width=True, margins=(0.4, 0.4, 0.75, 0.75)):
    ws.page_setup.paperSize = 9
    ws.page_setup.orientation = "portrait"
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_margins.left, ws.page_margins.right, ws.page_margins.top, ws.page_margins.bottom = margins
    ws.print_options.horizontalCentered = True
    ws.sheet_view.showGridLines = False


def rumus_bulat(mode, ref):
    if mode == "UP":
        return f"=CEILING({ref},1000)"
    if mode == "NONE":
        return f"={ref}"
    return f"=ROUND({ref},-3)"


# ======================================================================= KALKULASI
def sheet_kalkulasi(ws, trip: Trip, res: Result, st: Settings):
    a4(ws, margins=(0.24, 0.24, 0.6, 0.4))
    ws.page_setup.fitToHeight = 1
    for col, w in dict(A=3.4, B=2.3, C=31, D=1.9, E=17, F=3.7, G=5, H=30, I=22.4).items():
        ws.column_dimensions[col].width = w

    ws.merge_cells("A1:I1")
    setc(ws, "A1", "KALKULASI BIAYA PERJALANAN DINAS ", Font(name="Calibri", size=12, bold=True), A("center"))
    setc(ws, "A3", "Nama ", F()); setc(ws, "D3", ":", F()); setc(ws, "E3", trip.nama.upper(), F(11, True))
    setc(ws, "A4", "Tujuan", F()); setc(ws, "D4", ":", F())
    ws.merge_cells("E4:I4"); setc(ws, "E4", trip.tujuan_dokumen(), F(), A("left"))
    setc(ws, "A5", "Tgl / Lamanya", F()); setc(ws, "D5", ":", F())
    setc(ws, "E5", f"{rentang_tgl(trip.tgl_berangkat, trip.tgl_kembali)} / {trip.lama} HARI", F())

    r = 7
    i_cells = []
    no = [0]

    def nomor():
        no[0] += 1
        setc(ws, f"A{r}", no[0], F(), A("left"), "0")

    def hitung_row(rr, harga, qty, note=None):
        setc(ws, f"E{rr}", harga, F(), A("right"), NUM_RP)
        setc(ws, f"F{rr}", "x", F(), A("center"))
        setc(ws, f"G{rr}", qty, F(), A("center"), "0")
        if note:
            setc(ws, f"H{rr}", note, F(), A("left"))
        setc(ws, f"I{rr}", f"=E{rr}*G{rr}", F(), A("right"), NUM_RP)
        i_cells.append(f"I{rr}")

    def head(label, sub=False, colon=True, numbered=True):
        if numbered:
            nomor()
        if sub:
            setc(ws, f"B{r}", "-", F())
        setc(ws, f"C{r}", label, F())
        if colon:
            setc(ws, f"D{r}", ":", F())

    # 1. Tiket pesawat pergi / pulang
    for idx, kat in enumerate(("Tiket Pesawat Pergi", "Tiket Pesawat Pulang")):
        rows = res.items.get(kat) or [None]
        for j, ln in enumerate(rows):
            if j == 0:
                head(kat, sub=True, numbered=(idx == 0))
            if ln:
                setc(ws, f"E{r}", ln.label, F(), A("left"))
                setc(ws, f"I{r}", ln.jumlah if ln.qty == 1 else f"={ln.qty:g}*{ln.harga}", F(), A("right"), NUM_RP)
                i_cells.append(f"I{r}")
            r += 2

    # 2, 3. Kapal laut, bus/kereta
    for kat in ("Tiket Kapal Laut", "Tiket Bus / Kereta Api"):
        head(kat)
        r += 1
        rows = res.items.get(kat) or [None]
        for ln in rows:
            hitung_row(r, ln.harga if ln else 0, ln.qty if ln else 0, ln.label if ln else None)
            r += 1
        r += 1

    # 4. Iuran wajib
    rows = res.items.get("Iuran Wajib / Kontribusi") or [None]
    for j, ln in enumerate(rows):
        if j == 0:
            head("Iuran Wajib / Kontribusi", colon=False)
        hitung_row(r, ln.harga if ln else 0, ln.qty if ln else 0, ln.label if ln else None)
        r += 2

    # 5. Uang harian + representasi
    for j, ln in enumerate(res.harian):
        if j == 0:
            nomor()
        setc(ws, f"C{r}", ln.label, F()); setc(ws, f"D{r}", ":", F())
        hitung_row(r, ln.harga, ln.qty)
        r += 2
    if not res.harian:
        nomor(); setc(ws, f"C{r}", "Uang Harian", F()); setc(ws, f"D{r}", ":", F())
        hitung_row(r, 0, 0); r += 2
    for ln in res.representasi:
        setc(ws, f"C{r}", ln.label, F())
        hitung_row(r, ln.harga, ln.qty)
        r += 2

    # 6. Hotel
    hotels = res.hotel30 + res.hotel_riil
    if not hotels:
        nomor(); setc(ws, f"C{r}", "Biaya Hotel", F()); setc(ws, f"D{r}", ":", F())
        hitung_row(r, 0, 0); r += 2
    for j, ln in enumerate(hotels):
        if j == 0:
            nomor()
        setc(ws, f"C{r}", ln.label.split(" (30%")[0], F()); setc(ws, f"D{r}", ":", F())
        hitung_row(r, ln.harga, ln.qty, ln.note)
        r += 2

    # 7. Transit, 8. Sewa kendaraan
    for kat, extra in (("Biaya Transit", None), ("Sewa Kendaraan", "(Mobil/Kapal/Perahu)")):
        rows = res.items.get(kat) or [None]
        for j, ln in enumerate(rows):
            if j == 0:
                head(kat)
            hitung_row(r, ln.harga if ln else 0, ln.qty if ln else 0, ln.label if ln else None)
            r += 1
            if j == 0 and extra:
                setc(ws, f"C{r}", extra, F()); r += 1
            if j == len(rows) - 1:
                r += 1

    # 9. Biaya transportasi (nota + transportasi riil tanpa bukti)
    head("Biaya Transportasi")
    r += 1
    n9 = 0
    short = ["Taksi tempat kedudukan - bandara/terminal (PP)", "Taksi bandara/terminal - tempat tujuan (PP)",
             "Transportasi antar tempat tujuan"]
    for ln in res.items.get("Biaya Transportasi", []):
        hitung_row(r, ln.harga, ln.qty, ln.label); r += 1; n9 += 1
    for k, ln in enumerate(res.trans_riil):
        if ln.qty and ln.harga:
            setc(ws, f"C{r}", short[k], F(10), A("left", wrap=True))
            ws.row_dimensions[r].height = 27
            hitung_row(r, ln.harga, ln.qty); r += 1; n9 += 1
    if not n9:
        hitung_row(r, 0, 0); r += 1
    r += 1

    first, last_item = 7, r - 1
    setc(ws, f"G{r}", "Jumlah", F(), A("center"))
    setc(ws, f"I{r}", f"=SUM(I{first}:I{last_item})", F(11, True), A("right"), NUM_RP)
    ws[f"I{r}"].border = Border(right=THIN, top=THIN)
    jum = r
    r += 1
    setc(ws, f"G{r}", "Dibulatkan", F(), A("center"))
    setc(ws, f"I{r}", rumus_bulat(st.pembulatan, f"I{jum}"), F(11, True), A("right"), NUM_RP)
    dibulat = r
    r += 2
    setc(ws, f"C{r}", f"Untuk 1 Orang Selama {trip.lama} Hari", F())
    r += 1
    box(ws, 3, r, 1, 9)
    for c in range(1, 10):                                  # garis bawah blok identitas
        cell = ws.cell(5, c)
        cell.border = Border(left=cell.border.left, right=cell.border.right, bottom=THIN)
    r += 2

    ws.merge_cells(f"G{r}:I{r}")
    setc(ws, f"G{r}", f"{st.kota.title()}, {tgl_id(trip.tgl_dokumen)}", F(), A("center"))
    r += 1
    setc(ws, f"C{r}", "Mengetahui", F(), A("center"))
    r += 1
    ws.merge_cells(f"G{r}:I{r}")
    setc(ws, f"C{r}", st.sekwan_jabatan + ",", F(), A("center"))
    setc(ws, f"G{r}", st.pptk_jabatan + ",", F(), A("center"))
    r += 5
    ws.merge_cells(f"G{r}:I{r}")
    setc(ws, f"C{r}", st.sekwan_nama, F(11, u=True), A("center"))
    setc(ws, f"G{r}", st.pptk_nama, F(11, u=True), A("center"))
    r += 1
    ws.merge_cells(f"G{r}:I{r}")
    setc(ws, f"C{r}", f"NIP. {st.sekwan_nip}", F(), A("center"))
    setc(ws, f"G{r}", f"NIP. {st.pptk_nip}", F(), A("center"))
    ws.print_area = f"A1:I{r + 1}"
    return dibulat


# ======================================================================= DP RIIL
def sheet_dp_riil(ws, trip: Trip, res: Result, st: Settings):
    a4(ws, margins=(0.42, 0.3, 0.75, 0.75))
    for col, w in dict(A=9.1, B=44, C=3.7, D=5.3, E=4.0, F=4.4, G=11, H=3.6, I=12.5, J=5.3, K=13.5).items():
        ws.column_dimensions[col].width = w

    ws.merge_cells("A2:K2")
    setc(ws, "A2", "DAFTAR PENGELUARAN RILL", F(16, True, True), A("center"))
    setc(ws, "A5", "Yang bertanda tangan dibawah ini :", F(12))
    for r_, lab, val, bold in ((7, "Nama", trip.nama.upper(), True), (8, "NIP", trip.nip or "-", False),
                               (9, "Jabatan", trip.jabatan_dok.upper(), True)):
        setc(ws, f"A{r_}", lab, F(12), A(None))
        setc(ws, f"B{r_}", f":  {val}", F(12, bold), A(None))
    spd = trip.spd_no.strip() or "    /SPD.LD/Setwan"
    setc(ws, "A11", f"Berdasarkan Surat Perjalanan Dinas Nomor :  {spd}, tanggal {tgl_id(trip.spd_tgl)},", F(12))
    setc(ws, "A12", "dengan ini saya menyatakan dengan sesungguhnya bahwa:", F(12))
    setc(ws, "A14", 1, F(12), A("center"))
    for r_, t in ((14, "Biaya uang harian, uang representasi, dan biaya hotel/penginapan dibawah ini yang tidak "),
                  (15, "dapat diperoleh bukti-bukti pengeluarannya (dibayarkan secara lumpsum), meliputi :")):
        setc(ws, f"B{r_}", t, F(12))

    ws.merge_cells("B18:I18")
    for c, t in (("A18", "No"), ("B18", "Uraian"), ("K18", "Jumlah")):
        setc(ws, c, t, F(10, True), A("center"))
    for c in range(1, 12):
        ws.cell(18, c).border = Border(top=THIN, bottom=THIN, left=THIN if c in (1, 2, 10) else None,
                                       right=THIN if c in (1, 9, 10, 11) else None)
    r = 19
    k_cells = []

    def detail(rr, lab, qty, unit, harga, first_no=None, merge_lab=True):
        setc(ws, f"A{rr}", first_no, F(10), A("center"))
        if lab:
            setc(ws, f"B{rr}", lab, F(10), A("left"))
        setc(ws, f"C{rr}", qty, F(10), A("right"))
        setc(ws, f"D{rr}", unit, F(10), A("left"))
        setc(ws, f"E{rr}", "x", F(10), A("left"))
        setc(ws, f"F{rr}", "Rp", F(10), A("left"))
        setc(ws, f"G{rr}", harga, F(10), A("left"), NUM_ACC)
        setc(ws, f"H{rr}", "=", F(10), A("left"))
        setc(ws, f"I{rr}", f"=C{rr}*G{rr}", F(10), A("left"), NUM_ACC)
        setc(ws, f"J{rr}", "Rp.", F(10), A("left"))
        setc(ws, f"K{rr}", f"=I{rr}", F(10), A("left"), NUM_ACC)
        k_cells.append(f"K{rr}")
        ws.row_dimensions[rr].height = 18

    n = 0
    groups = [res.harian or [None], res.representasi or [None], res.hotel30 or [None]]
    for g_i, g in enumerate(groups):
        n = g_i + 1
        for j, ln in enumerate(g):
            if ln is None:
                lab = ["Uang Harian", "Uang Representasi", "Biaya Hotel (30%)"][g_i]
                detail(r, lab, 0, "hari", 0, n if j == 0 else None)
            else:
                detail(r, ln.label, ln.qty, ln.unit, ln.harga, n if j == 0 else None)
            r += 1
    # Catatan: biaya transportasi TIDAK dimasukkan di lembar ini - lihat lembar KALKULASI.
    # garis tabel
    for rr in range(19, r):
        for c in range(1, 12):
            cell = ws.cell(rr, c)
            cell.border = Border(left=THIN if c in (1, 2, 10) else None,
                                 right=THIN if c in (1, 9, 10, 11) else None)
    ws.merge_cells(f"B{r}:I{r}")
    setc(ws, f"B{r}", "Jumlah", F(11, True), A("center"))
    setc(ws, f"K{r}", f"=SUM({k_cells[0]}:{k_cells[-1]})" if k_cells else 0, F(11, True), A("left"), NUM_ACC)
    for c in range(1, 12):
        ws.cell(r, c).border = Border(top=THIN, bottom=THIN, left=THIN if c in (1, 2, 10) else None,
                                      right=THIN if c in (1, 9, 10, 11) else None)
    ws.row_dimensions[r].height = 18
    r += 2
    setc(ws, f"A{r}", 2, F(12), A("center"))
    for t in ("Jumlah uang tersebut pada angka 1 diatas benar-benar dikeluarkan untuk pelaksanaan perjalanan ",
              "dinas dimaksud dan apabila dikemudian hari terdapat kelebihan atas pembayaran, kami bersedia",
              f" untuk mengembalikan/menyetorkan kelebihan tersebut ke  Kas Daerah Pemerintah {st.kota.title() and 'Kota ' + st.kota.title()}."):
        setc(ws, f"B{r}", t, F(12)); r += 1
    r += 1
    setc(ws, f"B{r}", "Demikian Daftar Pengeluaran Riil ini dibuat dengan sebenarnya, untuk dipergunakan menurut", F(12)); r += 1
    setc(ws, f"B{r}", "perlunya.", F(12)); r += 3

    setc(ws, f"H{r}", f"{st.kota.title()},  {tgl_id(trip.tgl_dokumen)}", F(12, True)); r += 1
    ws.merge_cells(f"A{r}:D{r}"); ws.merge_cells(f"G{r}:K{r}")
    setc(ws, f"A{r}", "Mengetahui/Menyetujui", F(12, True), A("center"))
    setc(ws, f"G{r}", "Yang Melaksanakan Perjalanan Dinas,", F(12, True), A("center")); r += 1
    ws.merge_cells(f"A{r}:D{r}")
    setc(ws, f"A{r}", st.sekwan_jabatan, F(12, True), A("center")); r += 4
    ws.merge_cells(f"A{r}:D{r}"); ws.merge_cells(f"G{r}:K{r}")
    setc(ws, f"A{r}", st.sekwan_nama, F(12, True, True), A("center"))
    setc(ws, f"G{r}", trip.nama.upper(), F(12, True), A("center")); r += 1
    ws.merge_cells(f"A{r}:D{r}")
    setc(ws, f"A{r}", f"NIP. {st.sekwan_nip}", F(12, True), A("center"))
    ws.print_area = f"A1:K{r + 1}"


# ======================================================================= SP TANGGUNG JAWAB
def sheet_sp(ws, trip: Trip, st: Settings, ref_total: str):
    a4(ws, margins=(0.7, 0.7, 0.75, 0.75))
    ws.column_dimensions["I"].width = 17.7
    ws.column_dimensions["J"].width = 0.3
    ws.column_dimensions["K"].width = 13.4
    ws.merge_cells("A1:K1")
    setc(ws, "A1", "SURAT PERNYATAAN PERTANGGUNGJAWABAN", F(16, True, True), A("center"))
    setc(ws, "A3", "Yang bertanda tangan dibawah ini :", F(12))
    for r_, lab, val, bold in ((5, "Nama", trip.nama.upper(), True), (6, "NIP", trip.nip or "-", False),
                               (7, "Jabatan", trip.jabatan_dok.upper(), True)):
        setc(ws, f"A{r_}", lab, F(12), A(None))
        setc(ws, f"B{r_}", f":  {val}", F(12, bold), A(None))
    setc(ws, "A9", "Menyatakan bahwa saya bertanggung jawab penuh atas kebenararan Pelaksanaan", F(12))
    setc(ws, "A10", "Perjalanan Dinas sesuai dengan :", F(12))
    setc(ws, "B12", "SPT Nomor       :", F(12)); setc(ws, "D12", trip.spt_no, F(12))
    setc(ws, "H12", f"Tanggal  {tgl_id(trip.spt_tgl)}", F(12))
    setc(ws, "B13", "SPD Nomor      :", F(12)); setc(ws, "D13", trip.spd_no or "     /SPD/Setwan", F(12))
    setc(ws, "H13", f"Tanggal  {tgl_id(trip.spd_tgl)}", F(12))
    setc(ws, "B15", "Jumlah Dana    :", F(12))
    ws.merge_cells("D15:G15")
    setc(ws, "D15", f"={ref_total}", F(12, True), A("left"), '"Rp. "#,##0",-"')
    ws.merge_cells("A17:K17")
    setc(ws, "A17", "Dokumen pertanggungjawaban perjalanan dinas  disampaikan sesuai ketentuan yang", F(12), A("left"))
    setc(ws, "A18", "berlaku untuk keperluan administrasi dan keperluan pemeriksaan aparat pengawas.", F(12))
    setc(ws, "A21", "Demikian surat pernyataan ini dibuat dengan sebenarnya.", F(12))
    for r_, t, b in ((24, f"{st.kota.title()}, {tgl_id(trip.tgl_dokumen)}", False),
                     (25, "Yang Melaksanakan Perjalanan Dinas,", True), (30, trip.nama.upper(), True)):
        ws.merge_cells(f"F{r_}:I{r_}")
        setc(ws, f"F{r_}", t, F(12, b), A("center"))
    ws.print_area = "A1:K31"


def export_xlsx(trip: Trip, st: Settings, path: str):
    res = hitung(trip)
    wb = Workbook()
    ws_dp = wb.active
    ws_dp.title = "DP RIIL"
    ws_sp = wb.create_sheet("SP TANGGUNG JAWAB")
    ws_k = wb.create_sheet("KALKULASI")
    sheet_dp_riil(ws_dp, trip, res, st)
    dib = sheet_kalkulasi(ws_k, trip, res, st)
    sheet_sp(ws_sp, trip, st, f"KALKULASI!I{dib}")
    wb.calculation.fullCalcOnLoad = True
    wb.save(path)
    return res
