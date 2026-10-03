from PySide6.QtCore import Qt, Signal
from PySide6.QtWidgets import (QAbstractItemView, QComboBox, QDialog, QFileDialog, QFormLayout,
                               QHBoxLayout, QHeaderView, QLabel, QLineEdit, QMessageBox,
                               QPushButton, QTableWidget, QTableWidgetItem, QVBoxLayout,
                               QWidget, QDialogButtonBox)

from ..db import Database


class AnggotaDialog(QDialog):
    def __init__(self, parent=None, nama="", jabatan="ANGGOTA", nip="-"):
        super().__init__(parent)
        self.setWindowTitle("Data anggota")
        self.setMinimumWidth(440)
        f = QFormLayout(self)
        f.setContentsMargins(20, 20, 20, 16); f.setVerticalSpacing(10)
        self.ed_nama = QLineEdit(nama)
        self.cb_jab = QComboBox(); self.cb_jab.addItems(["KETUA", "WAKIL KETUA", "ANGGOTA"])
        self.cb_jab.setCurrentText(jabatan.upper() if jabatan.upper() in ("KETUA", "WAKIL KETUA") else "ANGGOTA")
        self.ed_nip = QLineEdit(nip)
        f.addRow("Nama lengkap (dengan gelar)", self.ed_nama)
        f.addRow("Jabatan", self.cb_jab)
        f.addRow("NIP", self.ed_nip)
        bb = QDialogButtonBox(QDialogButtonBox.Save | QDialogButtonBox.Cancel)
        bb.accepted.connect(self._ok); bb.rejected.connect(self.reject)
        f.addRow(bb)

    def _ok(self):
        if not self.ed_nama.text().strip():
            QMessageBox.warning(self, "Nama kosong", "Nama harus diisi."); return
        self.accept()

    def values(self):
        return self.ed_nama.text().strip().upper() if False else self.ed_nama.text().strip(), \
            self.cb_jab.currentText(), self.ed_nip.text().strip() or "-"


class AnggotaPage(QWidget):
    changed = Signal()

    def __init__(self, db: Database):
        super().__init__()
        self.db = db
        v = QVBoxLayout(self); v.setContentsMargins(28, 22, 28, 22); v.setSpacing(6)
        t = QLabel("Data Pimpinan & Anggota DPRD"); t.setObjectName("PageTitle")
        s = QLabel("Daftar ini muncul pada pilihan nama pelaksana. Jabatan menentukan tarif hotel (pimpinan / anggota).")
        s.setObjectName("PageSub")
        v.addWidget(t); v.addWidget(s); v.addSpacing(8)
        h = QHBoxLayout()
        self.search = QLineEdit(); self.search.setPlaceholderText("Cari nama…"); self.search.textChanged.connect(self._filter)
        b_add = QPushButton("＋ Tambah"); b_add.setProperty("kind", "primary"); b_add.clicked.connect(self.tambah)
        b_edit = QPushButton("Ubah"); b_edit.clicked.connect(self.ubah)
        b_del = QPushButton("Hapus"); b_del.clicked.connect(self.hapus)
        b_imp = QPushButton("Impor dari Excel…"); b_imp.clicked.connect(self.impor)
        h.addWidget(self.search, 1); h.addWidget(b_imp); h.addWidget(b_edit); h.addWidget(b_del); h.addWidget(b_add)
        v.addLayout(h)
        self.tb = QTableWidget(0, 4)
        self.tb.setHorizontalHeaderLabels(["No", "Nama", "Jabatan", "NIP"])
        self.tb.setSelectionBehavior(QAbstractItemView.SelectRows)
        self.tb.setEditTriggers(QAbstractItemView.NoEditTriggers)
        self.tb.setAlternatingRowColors(True); self.tb.verticalHeader().setVisible(False)
        self.tb.setShowGrid(False)
        hh = self.tb.horizontalHeader(); hh.setSectionResizeMode(1, QHeaderView.Stretch)
        hh.setSectionResizeMode(0, QHeaderView.ResizeToContents)
        self.tb.doubleClicked.connect(lambda *_: self.ubah())
        v.addWidget(self.tb, 1)
        self.load()

    def load(self):
        rows = self.db.daftar_anggota()
        self.tb.setRowCount(len(rows))
        for i, r in enumerate(rows):
            for j, val in enumerate((str(i + 1), r["nama"], r["jabatan"], r["nip"])):
                it = QTableWidgetItem(val)
                if j == 0:
                    it.setData(Qt.UserRole, r["id"])
                self.tb.setItem(i, j, it)
            self.tb.setRowHeight(i, 38)
        self._filter()

    def _filter(self):
        q = self.search.text().strip().lower()
        for i in range(self.tb.rowCount()):
            self.tb.setRowHidden(i, bool(q) and q not in self.tb.item(i, 1).text().lower())

    def _sel(self):
        r = self.tb.currentRow()
        if r < 0:
            QMessageBox.information(self, "Pilih data", "Pilih satu baris terlebih dahulu."); return None
        return (self.tb.item(r, 0).data(Qt.UserRole), self.tb.item(r, 1).text(),
                self.tb.item(r, 2).text(), self.tb.item(r, 3).text())

    def tambah(self):
        d = AnggotaDialog(self)
        if d.exec():
            self.db.tambah_anggota(*d.values()); self.load(); self.changed.emit()

    def ubah(self):
        s = self._sel()
        if not s:
            return
        d = AnggotaDialog(self, s[1], s[2], s[3])
        if d.exec():
            self.db.ubah_anggota(s[0], *d.values()); self.load(); self.changed.emit()

    def hapus(self):
        s = self._sel()
        if s and QMessageBox.question(self, "Hapus", f"Hapus {s[1]}?") == QMessageBox.Yes:
            self.db.hapus_anggota(s[0]); self.load(); self.changed.emit()

    def impor(self):
        path, _ = QFileDialog.getOpenFileName(self, "Pilih file Excel", "", "Excel (*.xlsx *.xlsm)")
        if not path:
            return
        try:
            rows = baca_excel_anggota(path)
        except Exception as e:                       # noqa
            QMessageBox.critical(self, "Gagal membaca file", str(e)); return
        n = self.db.impor_anggota(rows)
        self.load(); self.changed.emit()
        QMessageBox.information(self, "Impor selesai", f"{n} anggota baru ditambahkan ({len(rows) - n} sudah ada).")


def baca_excel_anggota(path):
    """Baca kolom NAMA / JABATAN (/ NIP) dari Excel, format seperti DATABASE_ANGGOTA_DPRD.xlsx."""
    from openpyxl import load_workbook
    ws = load_workbook(path, data_only=True).worksheets[0]
    col = {}
    hdr_row = None
    for r in range(1, min(ws.max_row, 10) + 1):
        for c in range(1, ws.max_column + 1):
            v = ws.cell(r, c).value
            if isinstance(v, str):
                k = v.replace(" ", "").upper()
                if k == "NAMA": col["nama"] = c; hdr_row = r
                elif k == "JABATAN": col["jabatan"] = c
                elif k == "NIP": col["nip"] = c
        if "nama" in col:
            break
    if "nama" not in col:
        raise ValueError("Kolom 'NAMA' tidak ditemukan pada 10 baris pertama.")
    out = []
    for r in range(hdr_row + 1, ws.max_row + 1):
        nama = ws.cell(r, col["nama"]).value
        if not nama or not str(nama).strip():
            continue
        jab = str(ws.cell(r, col["jabatan"]).value or "ANGGOTA").strip().upper() if "jabatan" in col else "ANGGOTA"
        nip = str(ws.cell(r, col["nip"]).value or "-").strip() if "nip" in col else "-"
        out.append((str(nama).strip(), jab, nip))
    return out
